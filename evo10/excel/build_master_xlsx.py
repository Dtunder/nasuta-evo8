import pandas as pd
import glob
import os
import numpy as np
from openpyxl import load_workbook
from openpyxl.styles import Font, PatternFill

def main():
    # 1. Read and Process CSVs
    csv_files = glob.glob('evo10/logs/multiseed_raw_*.csv')
    if not csv_files:
        print("No CSV files found.")
        return

    dfs = [pd.read_csv(f) for f in csv_files]
    df_raw = pd.concat(dfs, ignore_index=True)
    df_raw = df_raw.drop_duplicates(subset=['model', 'seed'])

    # 2. Compute Summary
    summary_data = []
    for model, group in df_raw.groupby('model'):
        n = len(group)
        max_rounds = group['rounds_survived'].max()
        mean_rounds = group['rounds_survived'].mean()
        std = group['rounds_survived'].std()
        ci95 = 1.96 * std / np.sqrt(n) if n > 1 else 0
        survival_rate_pct = (group['rounds_survived'] >= 30).mean() * 100
        mean_stability = group['stability'].mean()
        dominant_death_cause = group['death_cause'].mode()[0] if not group['death_cause'].mode().empty else None

        summary_data.append({
            'model': model,
            'n': n,
            'max_rounds': max_rounds,
            'mean_rounds': mean_rounds,
            'std': std,
            'ci95': ci95,
            'survival_rate_pct': survival_rate_pct,
            'mean_stability': mean_stability,
            'dominant_death_cause': dominant_death_cause
        })
    df_summary = pd.DataFrame(summary_data)

    # 3. Create Excel File
    output_path = 'evo10/excel/evo10_master.xlsx'
    os.makedirs(os.path.dirname(output_path), exist_ok=True)
    with pd.ExcelWriter(output_path, engine='openpyxl') as writer:
        df_raw.to_excel(writer, sheet_name='raw', index=False)
        df_summary.to_excel(writer, sheet_name='summary', index=False)

        # Notes sheet
        df_notes = pd.DataFrame({'Caveat': ["Sovereign = roher MCTS, erreicht max ~9 Runden, KEINE 30; 30-Runden-Claim stammte aus separater handcodierter Heuristik."]})
        df_notes.to_excel(writer, sheet_name='notes', index=False)

    # 4. Format Excel File
    wb = load_workbook(output_path)

    # Bold headers
    for sheet_name in wb.sheetnames:
        ws = wb[sheet_name]
        for cell in ws[1]:
            cell.font = Font(bold=True)

    # Color survival_rate_pct column in 'summary' sheet
    ws_summary = wb['summary']
    survival_col_idx = None
    for idx, cell in enumerate(ws_summary[1], 1):
        if cell.value == 'survival_rate_pct':
            survival_col_idx = idx
            break

    if survival_col_idx:
        fill = PatternFill(start_color="FFFFE0", end_color="FFFFE0", fill_type="solid") # Light yellow
        for row in range(2, ws_summary.max_row + 1):
            ws_summary.cell(row=row, column=survival_col_idx).fill = fill

    wb.save(output_path)
    print(f"Successfully created {output_path}")

if __name__ == "__main__":
    main()
